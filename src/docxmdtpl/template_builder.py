from __future__ import annotations

from pathlib import Path

from docx import Document as create_document
from docx.document import Document as WordDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


class BuildTemplate:
    def __init__(self, output_path: str | Path, max_depth: int = 5) -> None:
        if max_depth < 1:
            raise ValueError("max_depth precisa ser maior ou igual a 1.")

        self.output_path = Path(output_path).expanduser()
        self.max_depth = max_depth

    def build(self) -> Path:
        document = create_document()
        self._enable_update_fields(document)
        title_paragraph = document.add_paragraph("{{ metadata.title | default('') }}")
        self._try_apply_style(title_paragraph, "Title")
        self._add_section_loop(document, depth=1)
        self._add_template_prototypes(document)

        self.output_path.parent.mkdir(parents=True, exist_ok=True)
        document.save(self.output_path)
        return self.output_path.resolve()

    def _add_section_loop(self, document: WordDocument, depth: int) -> None:
        item_name = f"section_{depth}"
        parent_name = "sections" if depth == 1 else f"section_{depth - 1}.children"

        document.add_paragraph(f"{{%p for {item_name} in {parent_name} %}}")
        heading_paragraph = document.add_paragraph(f"{{{{p {item_name}.heading_subdoc }}}}")
        self._try_apply_style(heading_paragraph, f"Heading {depth}")
        content_paragraph = document.add_paragraph(f"{{{{p {item_name}.content_subdoc }}}}")
        self._try_apply_style(content_paragraph, "Normal")

        if depth < self.max_depth:
            self._add_section_loop(document, depth=depth + 1)

        document.add_paragraph("{%p endfor %}")

    def _add_template_prototypes(self, document: WordDocument) -> None:
        self._add_figure_caption_template(document)
        document.add_paragraph()
        self._add_quote_block_template(document)
        document.add_paragraph()
        self._add_code_block_template(document)
        document.add_paragraph()
        self._add_markdown_table_template(document)

    @staticmethod
    def _enable_update_fields(document: WordDocument) -> None:
        settings_element = document.settings._element
        update_fields = settings_element.first_child_found_in("w:updateFields")

        if update_fields is None:
            update_fields = OxmlElement("w:updateFields")
            settings_element.insert_element_before(update_fields, "w:hdrShapeDefaults")

        update_fields.set(qn("w:val"), "true")

    @staticmethod
    def _try_apply_style(paragraph: object, style_name: str) -> None:
        try:
            paragraph.style = style_name
        except (KeyError, ValueError):
            pass

    def _add_figure_caption_template(self, document: WordDocument) -> None:
        caption_paragraph = document.add_paragraph()
        self._try_apply_style(caption_paragraph, "Caption")
        caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption_paragraph.add_run("Figura ")
        self._append_sequence_field(caption_paragraph, sequence_name="Figura", current_value=1)
        caption_paragraph.add_run(" - {{ figure_caption }}")

    def _add_quote_block_template(self, document: WordDocument) -> None:
        quote_table = document.add_table(rows=1, cols=1)
        self._try_apply_table_style(quote_table, "Table Grid")
        quote_paragraph = quote_table.cell(0, 0).paragraphs[0]
        quote_paragraph.text = "{{ quote_block }}"
        self._try_apply_style(quote_paragraph, "Quote")

    def _add_code_block_template(self, document: WordDocument) -> None:
        code_table = document.add_table(rows=1, cols=1)
        self._try_apply_table_style(code_table, "Table Grid")
        code_paragraph = code_table.cell(0, 0).paragraphs[0]
        code_paragraph.text = "{{ code_block }}"
        if code_paragraph.runs:
            code_paragraph.runs[0].font.name = "Consolas"

    def _add_markdown_table_template(self, document: WordDocument) -> None:
        table = document.add_table(rows=3, cols=2)
        self._try_apply_table_style(table, "Table Grid")

        header_row = table.rows[0].cells
        header_row[0].text = "{{ table_header_1 }}"
        header_row[1].text = "{{ table_header_2 }}"
        for cell in header_row:
            if cell.paragraphs[0].runs:
                cell.paragraphs[0].runs[0].bold = True

        odd_row = table.rows[1].cells
        odd_row[0].text = "{{ table_odd_1 }}"
        odd_row[1].text = "{{ table_odd_2 }}"

        even_row = table.rows[2].cells
        even_row[0].text = "{{ table_even_1 }}"
        even_row[1].text = "{{ table_even_2 }}"

    @staticmethod
    def _append_sequence_field(paragraph: object, *, sequence_name: str, current_value: int) -> None:
        begin_run = paragraph.add_run()
        begin = OxmlElement("w:fldChar")
        begin.set(qn("w:fldCharType"), "begin")
        begin_run._r.append(begin)

        instruction_run = paragraph.add_run()
        instruction = OxmlElement("w:instrText")
        instruction.set(qn("xml:space"), "preserve")
        instruction.text = f" SEQ {sequence_name} \\* ARABIC "
        instruction_run._r.append(instruction)

        separate_run = paragraph.add_run()
        separate = OxmlElement("w:fldChar")
        separate.set(qn("w:fldCharType"), "separate")
        separate_run._r.append(separate)

        paragraph.add_run(str(current_value))

        end_run = paragraph.add_run()
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        end_run._r.append(end)

    @staticmethod
    def _try_apply_table_style(table: object, style_name: str) -> None:
        try:
            table.style = style_name
        except (KeyError, ValueError):
            pass
