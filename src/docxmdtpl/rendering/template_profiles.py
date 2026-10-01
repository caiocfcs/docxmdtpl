from __future__ import annotations

from copy import deepcopy
import re
from dataclasses import dataclass, field
from typing import Any

from docx.oxml.ns import qn
from docx.shared import RGBColor

from docxmdtpl.errors import TemplateBindingError


LOOP_START_PATTERN = re.compile(
    r"^\{\%p\s+for\s+(?P<variable>[A-Za-z_][A-Za-z0-9_]*)\s+in\s+(?P<source>.+?)\s*\%\}$"
)
LOOP_END_PATTERN = re.compile(r"^\{\%p\s+endfor\s+\%\}$")
PLACEHOLDER_PATTERN = re.compile(
    r"^\{\{\s*p\s+(?P<variable>[A-Za-z_][A-Za-z0-9_]*)\.(?P<field>heading_subdoc|content_subdoc)\s*\}\}$"
)
CHILDREN_SOURCE_PATTERN = re.compile(r"^(?P<parent>[A-Za-z_][A-Za-z0-9_]*)\.children$")
FIGURE_CAPTION_PATTERN = re.compile(r"\{\{\s*figure_caption\s*\}\}")
QUOTE_BLOCK_PATTERN = re.compile(r"^\{\{\s*quote_block\s*\}\}$")
CODE_BLOCK_PATTERN = re.compile(r"^\{\{\s*code_block\s*\}\}$")


@dataclass(slots=True)
class RunFormatProfile:
    style_name: str | None = None
    run_properties_xml: Any | None = None
    font_name: str | None = None
    font_size: Any | None = None
    bold: bool | None = None
    italic: bool | None = None
    underline: Any | None = None
    strike: bool | None = None
    subscript: bool | None = None
    superscript: bool | None = None
    all_caps: bool | None = None
    color_rgb: RGBColor | None = None


@dataclass(slots=True)
class ParagraphFormatProfile:
    style_name: str | None = None
    alignment: Any | None = None
    left_indent: Any | None = None
    right_indent: Any | None = None
    first_line_indent: Any | None = None
    space_before: Any | None = None
    space_after: Any | None = None
    line_spacing: Any | None = None
    line_spacing_rule: Any | None = None
    keep_together: bool | None = None
    keep_with_next: bool | None = None
    page_break_before: bool | None = None
    widow_control: bool | None = None


@dataclass(slots=True)
class PlaceholderProfile:
    paragraph: ParagraphFormatProfile = field(default_factory=ParagraphFormatProfile)
    run: RunFormatProfile = field(default_factory=RunFormatProfile)


@dataclass(slots=True)
class FigureCaptionTemplate:
    paragraph_xml: Any


@dataclass(slots=True)
class BlockContainerTemplate:
    table_properties_xml: Any | None = None
    grid_column_widths: list[Any] = field(default_factory=list)
    row_properties_xml: Any | None = None
    cell_properties_xml: Any | None = None
    placeholder_profile: PlaceholderProfile | None = None


@dataclass(slots=True)
class TableCellTemplate:
    cell_properties_xml: Any | None = None
    paragraph_profile: PlaceholderProfile | None = None


@dataclass(slots=True)
class TableRowTemplate:
    row_properties_xml: Any | None = None
    cell_templates: list[TableCellTemplate] = field(default_factory=list)

    def resolve_cell(self, column_index: int) -> TableCellTemplate | None:
        if not self.cell_templates:
            return None
        if column_index < len(self.cell_templates):
            return self.cell_templates[column_index]
        return self.cell_templates[-1]


@dataclass(slots=True)
class MarkdownTableTemplate:
    table_properties_xml: Any | None = None
    grid_column_widths: list[Any] = field(default_factory=list)
    header_row: TableRowTemplate = field(default_factory=TableRowTemplate)
    odd_row: TableRowTemplate = field(default_factory=TableRowTemplate)
    even_row: TableRowTemplate = field(default_factory=TableRowTemplate)


@dataclass(slots=True)
class TemplateRenderProfiles:
    heading_profiles: dict[int, PlaceholderProfile] = field(default_factory=dict)
    content_profiles: dict[int, PlaceholderProfile] = field(default_factory=dict)
    figure_caption_templates: dict[int, FigureCaptionTemplate] = field(default_factory=dict)
    quote_templates: dict[int, BlockContainerTemplate] = field(default_factory=dict)
    code_templates: dict[int, BlockContainerTemplate] = field(default_factory=dict)
    table_templates: dict[int, MarkdownTableTemplate] = field(default_factory=dict)

    def resolve_heading(self, level: int) -> PlaceholderProfile | None:
        return self._resolve(self.heading_profiles, level)

    def resolve_content(self, level: int) -> PlaceholderProfile | None:
        return self._resolve(self.content_profiles, level)

    def resolve_figure_caption(self, level: int) -> FigureCaptionTemplate | None:
        return self._resolve(self.figure_caption_templates, level)

    def resolve_quote(self, level: int) -> BlockContainerTemplate | None:
        return self._resolve(self.quote_templates, level)

    def resolve_code(self, level: int) -> BlockContainerTemplate | None:
        return self._resolve(self.code_templates, level)

    def resolve_table(self, level: int) -> MarkdownTableTemplate | None:
        return self._resolve(self.table_templates, level)

    @staticmethod
    def _resolve(profiles: dict[int, PlaceholderProfile], level: int) -> PlaceholderProfile | None:
        if not profiles:
            return None
        if level in profiles:
            return profiles[level]

        shallower_levels = [depth for depth in profiles if depth <= level]
        if shallower_levels:
            return profiles[max(shallower_levels)]

        return profiles[min(profiles)]


@dataclass(slots=True)
class _LoopFrame:
    variable: str
    depth: int


class TemplateProfileExtractor:
    def extract(self, template: Any) -> TemplateRenderProfiles:
        document = template.get_docx()
        profiles = TemplateRenderProfiles()
        loop_stack: list[_LoopFrame] = []
        variable_depths: dict[str, int] = {}
        paragraph_map = {paragraph._p: paragraph for paragraph in document.paragraphs}
        table_map = {table._tbl: table for table in document.tables}

        for element in list(document._element.body.iterchildren()):
            paragraph = paragraph_map.get(element)
            if paragraph is not None:
                text = paragraph.text.strip()
                if not text:
                    continue

                loop_start = LOOP_START_PATTERN.fullmatch(text)
                if loop_start is not None:
                    variable = loop_start.group("variable")
                    source = loop_start.group("source").strip()
                    depth = self._resolve_loop_depth(source, loop_stack, variable_depths)
                    loop_stack.append(_LoopFrame(variable=variable, depth=depth))
                    variable_depths[variable] = depth
                    continue

                if LOOP_END_PATTERN.fullmatch(text) is not None:
                    if loop_stack:
                        frame = loop_stack.pop()
                        variable_depths.pop(frame.variable, None)
                    continue

                if FIGURE_CAPTION_PATTERN.search(text) is not None:
                    depth = loop_stack[-1].depth if loop_stack else 1
                    profiles.figure_caption_templates.setdefault(
                        depth,
                        FigureCaptionTemplate(paragraph_xml=deepcopy(paragraph._p)),
                    )
                    self._remove_paragraph(paragraph)
                    continue

                placeholder = PLACEHOLDER_PATTERN.fullmatch(text)
                if placeholder is None:
                    continue

                variable = placeholder.group("variable")
                field_name = placeholder.group("field")
                depth = variable_depths.get(variable, loop_stack[-1].depth if loop_stack else 1)
                profile = self._build_profile(paragraph)
                field_profiles = profiles.heading_profiles if field_name == "heading_subdoc" else profiles.content_profiles
                field_profiles.setdefault(depth, profile)
                continue

            table = table_map.get(element)
            if table is None:
                continue

            depth = loop_stack[-1].depth if loop_stack else 1
            quote_template = self._extract_block_container_template(table, QUOTE_BLOCK_PATTERN)
            if quote_template is not None:
                profiles.quote_templates.setdefault(depth, quote_template)
                self._remove_table(table)
                continue

            code_template = self._extract_block_container_template(table, CODE_BLOCK_PATTERN)
            if code_template is not None:
                profiles.code_templates.setdefault(depth, code_template)
                self._remove_table(table)
                continue

            table_template = self._extract_markdown_table_template(table)
            if table_template is not None:
                profiles.table_templates.setdefault(depth, table_template)
                self._remove_table(table)

        return profiles

    @staticmethod
    def _resolve_loop_depth(
        source: str,
        loop_stack: list[_LoopFrame],
        variable_depths: dict[str, int],
    ) -> int:
        if source == "sections":
            return 1

        children_source = CHILDREN_SOURCE_PATTERN.fullmatch(source)
        if children_source is not None:
            parent = children_source.group("parent")
            if parent in variable_depths:
                return variable_depths[parent] + 1

        return len(loop_stack) + 1

    def _build_profile(self, paragraph: Any) -> PlaceholderProfile:
        paragraph_style = getattr(paragraph, "style", None)
        paragraph_style_format = getattr(paragraph_style, "paragraph_format", None)
        paragraph_profile = ParagraphFormatProfile(
            style_name=getattr(paragraph_style, "name", None),
            alignment=_first_defined(
                paragraph.alignment,
                getattr(paragraph_style_format, "alignment", None),
            ),
            left_indent=_first_defined(
                paragraph.paragraph_format.left_indent,
                getattr(paragraph_style_format, "left_indent", None),
            ),
            right_indent=_first_defined(
                paragraph.paragraph_format.right_indent,
                getattr(paragraph_style_format, "right_indent", None),
            ),
            first_line_indent=_first_defined(
                paragraph.paragraph_format.first_line_indent,
                getattr(paragraph_style_format, "first_line_indent", None),
            ),
            space_before=_first_defined(
                paragraph.paragraph_format.space_before,
                getattr(paragraph_style_format, "space_before", None),
            ),
            space_after=_first_defined(
                paragraph.paragraph_format.space_after,
                getattr(paragraph_style_format, "space_after", None),
            ),
            line_spacing=_first_defined(
                paragraph.paragraph_format.line_spacing,
                getattr(paragraph_style_format, "line_spacing", None),
            ),
            line_spacing_rule=_first_defined(
                paragraph.paragraph_format.line_spacing_rule,
                getattr(paragraph_style_format, "line_spacing_rule", None),
            ),
            keep_together=_first_defined(
                paragraph.paragraph_format.keep_together,
                getattr(paragraph_style_format, "keep_together", None),
            ),
            keep_with_next=_first_defined(
                paragraph.paragraph_format.keep_with_next,
                getattr(paragraph_style_format, "keep_with_next", None),
            ),
            page_break_before=_first_defined(
                paragraph.paragraph_format.page_break_before,
                getattr(paragraph_style_format, "page_break_before", None),
            ),
            widow_control=_first_defined(
                paragraph.paragraph_format.widow_control,
                getattr(paragraph_style_format, "widow_control", None),
            ),
        )

        reference_run = self._select_reference_run(paragraph)
        if reference_run is None:
            return PlaceholderProfile(paragraph=paragraph_profile)

        run_style = getattr(reference_run, "style", None)
        run_style_font = getattr(run_style, "font", None)
        paragraph_style_font = getattr(paragraph_style, "font", None)
        font = reference_run.font
        run_profile = RunFormatProfile(
            style_name=self._resolve_run_style_name(run_style),
            run_properties_xml=deepcopy(getattr(reference_run._r, "rPr", None)),
            font_name=_first_defined(
                font.name,
                getattr(run_style_font, "name", None),
                getattr(paragraph_style_font, "name", None),
            ),
            font_size=_first_defined(
                font.size,
                getattr(run_style_font, "size", None),
                getattr(paragraph_style_font, "size", None),
            ),
            bold=_first_defined(
                reference_run.bold,
                getattr(run_style_font, "bold", None),
                getattr(paragraph_style_font, "bold", None),
            ),
            italic=_first_defined(
                reference_run.italic,
                getattr(run_style_font, "italic", None),
                getattr(paragraph_style_font, "italic", None),
            ),
            underline=_first_defined(
                reference_run.underline,
                getattr(run_style_font, "underline", None),
                getattr(paragraph_style_font, "underline", None),
            ),
            strike=_first_defined(
                font.strike,
                getattr(run_style_font, "strike", None),
                getattr(paragraph_style_font, "strike", None),
            ),
            subscript=_first_defined(
                font.subscript,
                getattr(run_style_font, "subscript", None),
                getattr(paragraph_style_font, "subscript", None),
            ),
            superscript=_first_defined(
                font.superscript,
                getattr(run_style_font, "superscript", None),
                getattr(paragraph_style_font, "superscript", None),
            ),
            all_caps=_first_defined(
                font.all_caps,
                getattr(run_style_font, "all_caps", None),
                getattr(paragraph_style_font, "all_caps", None),
            ),
            color_rgb=self._resolve_color_rgb(
                reference_run,
                run_style_font,
                paragraph_style_font,
            ),
        )
        return PlaceholderProfile(paragraph=paragraph_profile, run=run_profile)

    @staticmethod
    def _select_reference_run(paragraph: Any) -> Any | None:
        for run in paragraph.runs:
            if run.text and run.text.strip():
                return run
        if paragraph.runs:
            return paragraph.runs[0]
        return None

    @staticmethod
    def _resolve_run_style_name(run_style: Any) -> str | None:
        style_name = getattr(run_style, "name", None)
        if style_name in {None, "Default Paragraph Font"}:
            return None
        return style_name

    @staticmethod
    def _resolve_color_rgb(run: Any, run_style_font: Any, paragraph_style_font: Any) -> RGBColor | None:
        run_properties = getattr(run._r, "rPr", None)
        if run_properties is not None and run_properties.find(qn("w:color")) is not None:
            return run.font.color.rgb
        return _first_defined(
            run.font.color.rgb,
            getattr(getattr(run_style_font, "color", None), "rgb", None),
            getattr(getattr(paragraph_style_font, "color", None), "rgb", None),
        )

    @staticmethod
    def _remove_paragraph(paragraph: Any) -> None:
        element = paragraph._element
        parent = element.getparent()
        if parent is not None:
            parent.remove(element)

    def _extract_block_container_template(
        self,
        table: Any,
        marker_pattern: re.Pattern[str],
    ) -> BlockContainerTemplate | None:
        matching_cells = []

        for row in table.rows:
            for cell in row.cells:
                cell_text = cell.text.strip()
                if marker_pattern.fullmatch(cell_text):
                    matching_cells.append(cell)
                    continue
                if marker_pattern.search(cell_text):
                    raise TemplateBindingError(
                        "O marcador de prototipo precisa ocupar sozinho a unica celula da tabela 1x1.",
                        template_path=None,
                    )

        if not matching_cells:
            return None

        if len(table.rows) != 1 or len(table.columns) != 1 or len(matching_cells) != 1:
            raise TemplateBindingError(
                "Os prototipos de quote_block e code_block precisam usar uma tabela 1x1.",
                template_path=None,
            )

        placeholder_paragraph = matching_cells[0].paragraphs[0] if matching_cells[0].paragraphs else None
        placeholder_profile = self._build_profile(placeholder_paragraph) if placeholder_paragraph is not None else None
        return BlockContainerTemplate(
            table_properties_xml=deepcopy(getattr(table._tbl, "tblPr", None)),
            grid_column_widths=[deepcopy(grid_col.w) for grid_col in table._tbl.tblGrid.gridCol_lst],
            row_properties_xml=deepcopy(getattr(table.rows[0]._tr, "trPr", None)),
            cell_properties_xml=deepcopy(getattr(matching_cells[0]._tc, "tcPr", None)),
            placeholder_profile=placeholder_profile,
        )

    @staticmethod
    def _remove_table(table: Any) -> None:
        element = table._tbl
        parent = element.getparent()
        if parent is not None:
            parent.remove(element)

    def _extract_markdown_table_template(self, table: Any) -> MarkdownTableTemplate | None:
        expected_rows = [
            ["{{ table_header_1 }}", "{{ table_header_2 }}"],
            ["{{ table_odd_1 }}", "{{ table_odd_2 }}"],
            ["{{ table_even_1 }}", "{{ table_even_2 }}"],
        ]
        actual_rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]

        if actual_rows != expected_rows:
            if any("{{ table_" in cell_text for row in actual_rows for cell_text in row):
                raise TemplateBindingError(
                    "O prototipo de tabela precisa ter 3 linhas x 2 colunas com os marcadores esperados.",
                    template_path=None,
                )
            return None

        return MarkdownTableTemplate(
            table_properties_xml=deepcopy(getattr(table._tbl, "tblPr", None)),
            grid_column_widths=[deepcopy(grid_col.w) for grid_col in table._tbl.tblGrid.gridCol_lst],
            header_row=self._build_table_row_template(table.rows[0]),
            odd_row=self._build_table_row_template(table.rows[1]),
            even_row=self._build_table_row_template(table.rows[2]),
        )

    def _build_table_row_template(self, row: Any) -> TableRowTemplate:
        return TableRowTemplate(
            row_properties_xml=deepcopy(getattr(row._tr, "trPr", None)),
            cell_templates=[self._build_table_cell_template(cell) for cell in row.cells],
        )

    def _build_table_cell_template(self, cell: Any) -> TableCellTemplate:
        placeholder_paragraph = cell.paragraphs[0] if cell.paragraphs else None
        paragraph_profile = self._build_profile(placeholder_paragraph) if placeholder_paragraph is not None else None
        return TableCellTemplate(
            cell_properties_xml=deepcopy(getattr(cell._tc, "tcPr", None)),
            paragraph_profile=paragraph_profile,
        )


def apply_paragraph_profile(
    paragraph: Any,
    profile: PlaceholderProfile | None,
    *,
    preserve_style: bool,
    preserve_indents: bool,
) -> None:
    if profile is None:
        return

    paragraph_profile = profile.paragraph

    if not preserve_style and paragraph_profile.style_name is not None:
        try:
            paragraph.style = paragraph_profile.style_name
        except (KeyError, ValueError):
            pass

    if paragraph_profile.alignment is not None:
        paragraph.alignment = paragraph_profile.alignment

    paragraph_format = paragraph.paragraph_format

    if not preserve_indents:
        if paragraph_profile.left_indent is not None:
            paragraph_format.left_indent = paragraph_profile.left_indent
        if paragraph_profile.right_indent is not None:
            paragraph_format.right_indent = paragraph_profile.right_indent
        if paragraph_profile.first_line_indent is not None:
            paragraph_format.first_line_indent = paragraph_profile.first_line_indent

    if paragraph_profile.space_before is not None:
        paragraph_format.space_before = paragraph_profile.space_before
    if paragraph_profile.space_after is not None:
        paragraph_format.space_after = paragraph_profile.space_after
    if paragraph_profile.line_spacing is not None:
        paragraph_format.line_spacing = paragraph_profile.line_spacing
    if paragraph_profile.line_spacing_rule is not None:
        paragraph_format.line_spacing_rule = paragraph_profile.line_spacing_rule
    if paragraph_profile.keep_together is not None:
        paragraph_format.keep_together = paragraph_profile.keep_together
    if paragraph_profile.keep_with_next is not None:
        paragraph_format.keep_with_next = paragraph_profile.keep_with_next
    if paragraph_profile.page_break_before is not None:
        paragraph_format.page_break_before = paragraph_profile.page_break_before
    if paragraph_profile.widow_control is not None:
        paragraph_format.widow_control = paragraph_profile.widow_control


def apply_run_profile(run: Any, profile: PlaceholderProfile | None) -> None:
    if profile is None:
        return

    run_profile = profile.run

    if run_profile.run_properties_xml is not None:
        _replace_run_properties(run, run_profile.run_properties_xml)

    if run_profile.style_name is not None:
        try:
            run.style = run_profile.style_name
        except (KeyError, ValueError):
            pass

    font = run.font

    if run_profile.font_name is not None:
        font.name = run_profile.font_name
    if run_profile.font_size is not None:
        font.size = run_profile.font_size
    if run_profile.bold is not None:
        run.bold = run_profile.bold
    if run_profile.italic is not None:
        run.italic = run_profile.italic
    if run_profile.underline is not None:
        run.underline = run_profile.underline
    if run_profile.strike is not None:
        font.strike = run_profile.strike
    if run_profile.subscript is not None:
        font.subscript = run_profile.subscript
    if run_profile.superscript is not None:
        font.superscript = run_profile.superscript
    if run_profile.all_caps is not None:
        font.all_caps = run_profile.all_caps
    if run_profile.color_rgb is not None:
        font.color.rgb = run_profile.color_rgb


def _replace_run_properties(run: Any, run_properties_xml: Any) -> None:
    existing_run_properties = getattr(run._r, "rPr", None)
    if existing_run_properties is not None:
        run._r.remove(existing_run_properties)
    run._r.insert(0, deepcopy(run_properties_xml))


def _first_defined(*values: Any) -> Any | None:
    for value in values:
        if value is not None:
            return value
    return None
