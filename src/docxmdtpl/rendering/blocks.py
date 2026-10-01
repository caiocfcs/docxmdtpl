from __future__ import annotations

from copy import deepcopy
from typing import Any

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Inches, Pt, RGBColor
from docx.table import Table
from docx.text.run import Run

from docxmdtpl.errors import TemplateBindingError
from docxmdtpl.errors import UnsupportedMarkdownError
from docxmdtpl.models import (
    BlockLike,
    BoldSpan,
    BoldItalicSpan,
    CodeBlock,
    CodeSpan,
    HorizontalRuleBlock,
    ImageBlock,
    InlineSpan,
    ItalicSpan,
    LinkSpan,
    ListBlock,
    ListItem,
    ParagraphBlock,
    QuoteBlock,
    SectionNode,
    SpacerBlock,
    TableBlock,
    TableCell,
    TextSpan,
)
from docxmdtpl.rendering.image_resolver import ImageResolver
from docxmdtpl.rendering.template_profiles import BlockContainerTemplate
from docxmdtpl.rendering.template_profiles import FIGURE_CAPTION_PATTERN
from docxmdtpl.rendering.template_profiles import FigureCaptionTemplate
from docxmdtpl.rendering.template_profiles import MarkdownTableTemplate
from docxmdtpl.rendering.template_profiles import PlaceholderProfile
from docxmdtpl.rendering.template_profiles import TableCellTemplate
from docxmdtpl.rendering.template_profiles import TableRowTemplate
from docxmdtpl.rendering.template_profiles import apply_paragraph_profile
from docxmdtpl.rendering.template_profiles import apply_run_profile


class ContentRenderer:
    def __init__(self, image_resolver: ImageResolver | None = None) -> None:
        self._image_resolver = image_resolver or ImageResolver()
        self._figure_sequence = 0

    def reset(self) -> None:
        self._figure_sequence = 0

    def render_content(
        self,
        section: SectionNode,
        tpl: object,
        path_imagens: str | None,
        source_dir: str | None,
        placeholder_profile: PlaceholderProfile | None = None,
        figure_caption_template: FigureCaptionTemplate | None = None,
        quote_block_template: BlockContainerTemplate | None = None,
        code_block_template: BlockContainerTemplate | None = None,
        markdown_table_template: MarkdownTableTemplate | None = None,
    ) -> object:
        subdoc = tpl.new_subdoc()

        for block in section.blocks:
            self._render_block(
                subdoc=subdoc,
                block=block,
                path_imagens=path_imagens,
                source_dir=source_dir,
                quote_depth=0,
                placeholder_profile=placeholder_profile,
                figure_caption_template=figure_caption_template,
                quote_block_template=quote_block_template,
                code_block_template=code_block_template,
                markdown_table_template=markdown_table_template,
            )

        return subdoc

    def _render_block(
        self,
        *,
        subdoc: Any,
        block: BlockLike,
        path_imagens: str | None,
        source_dir: str | None,
        quote_depth: int,
        placeholder_profile: PlaceholderProfile | None,
        figure_caption_template: FigureCaptionTemplate | None,
        quote_block_template: BlockContainerTemplate | None,
        code_block_template: BlockContainerTemplate | None,
        markdown_table_template: MarkdownTableTemplate | None,
    ) -> None:
        if isinstance(block, ParagraphBlock):
            paragraph_style = "Quote" if quote_depth > 0 else self._resolve_base_paragraph_style(placeholder_profile)
            paragraph = self._add_paragraph(subdoc, paragraph_style)
            apply_paragraph_profile(
                paragraph,
                placeholder_profile,
                preserve_style=quote_depth > 0,
                preserve_indents=quote_depth > 0,
            )
            self._add_inline_runs(paragraph, block.spans, placeholder_profile=placeholder_profile)
            return

        if isinstance(block, ListBlock):
            self._render_list(
                subdoc,
                block,
                depth=1,
                path_imagens=path_imagens,
                source_dir=source_dir,
                quote_depth=quote_depth,
                placeholder_profile=placeholder_profile,
                figure_caption_template=figure_caption_template,
                quote_block_template=quote_block_template,
                code_block_template=code_block_template,
                markdown_table_template=markdown_table_template,
            )
            return

        if isinstance(block, TableBlock):
            self._render_table(
                subdoc,
                block,
                path_imagens=path_imagens,
                source_dir=source_dir,
                placeholder_profile=placeholder_profile,
                figure_caption_template=figure_caption_template,
                markdown_table_template=markdown_table_template,
            )
            return

        if isinstance(block, QuoteBlock):
            self._render_quote(
                subdoc,
                block,
                path_imagens=path_imagens,
                source_dir=source_dir,
                quote_depth=quote_depth + 1,
                placeholder_profile=placeholder_profile,
                figure_caption_template=figure_caption_template,
                quote_block_template=quote_block_template,
                code_block_template=code_block_template,
                markdown_table_template=markdown_table_template,
            )
            return

        if isinstance(block, CodeBlock):
            self._render_code_block(
                subdoc,
                block,
                quote_depth=quote_depth,
                placeholder_profile=placeholder_profile,
                path_imagens=path_imagens,
                source_dir=source_dir,
                figure_caption_template=figure_caption_template,
                quote_block_template=quote_block_template,
                code_block_template=code_block_template,
                markdown_table_template=markdown_table_template,
            )
            return

        if isinstance(block, ImageBlock):
            self._render_image(
                subdoc,
                block,
                path_imagens,
                source_dir,
                placeholder_profile=placeholder_profile,
                figure_caption_template=figure_caption_template,
            )
            return

        if isinstance(block, HorizontalRuleBlock):
            self._render_horizontal_rule(subdoc, quote_depth=quote_depth, placeholder_profile=placeholder_profile)
            return

        if isinstance(block, SpacerBlock):
            self._render_spacer(
                subdoc,
                line_breaks=block.line_breaks,
                quote_depth=quote_depth,
                placeholder_profile=placeholder_profile,
            )
            return

        raise UnsupportedMarkdownError(
            "Bloco fora da cobertura do renderer da V1.",
            block_type=type(block).__name__,
        )

    def _render_list(
        self,
        subdoc: Any,
        list_block: ListBlock,
        depth: int,
        *,
        path_imagens: str | None,
        source_dir: str | None,
        quote_depth: int,
        placeholder_profile: PlaceholderProfile | None,
        figure_caption_template: FigureCaptionTemplate | None,
        quote_block_template: BlockContainerTemplate | None,
        code_block_template: BlockContainerTemplate | None,
        markdown_table_template: MarkdownTableTemplate | None,
    ) -> None:
        base_style = "List Number" if list_block.ordered else "List Bullet"
        nested_style = f"{base_style} {depth}" if depth > 1 else base_style

        for item in list_block.items:
            paragraph = self._add_paragraph(subdoc, nested_style, fallback_style=base_style)
            apply_paragraph_profile(
                paragraph,
                placeholder_profile,
                preserve_style=True,
                preserve_indents=True,
            )
            self._add_inline_runs(paragraph, item.spans, placeholder_profile=placeholder_profile)

            for child_block in item.children:
                if isinstance(child_block, ListBlock):
                    self._render_list(
                        subdoc,
                        child_block,
                        depth=depth + 1,
                        path_imagens=path_imagens,
                        source_dir=source_dir,
                        quote_depth=quote_depth,
                        placeholder_profile=placeholder_profile,
                        figure_caption_template=figure_caption_template,
                        quote_block_template=quote_block_template,
                        code_block_template=code_block_template,
                        markdown_table_template=markdown_table_template,
                    )
                    continue

                self._render_block(
                    subdoc=subdoc,
                    block=child_block,
                    path_imagens=path_imagens,
                    source_dir=source_dir,
                    quote_depth=quote_depth,
                    placeholder_profile=placeholder_profile,
                    figure_caption_template=figure_caption_template,
                    quote_block_template=quote_block_template,
                    code_block_template=code_block_template,
                    markdown_table_template=markdown_table_template,
                )

    def _render_table(
        self,
        subdoc: Any,
        table_block: TableBlock,
        *,
        path_imagens: str | None,
        source_dir: str | None,
        placeholder_profile: PlaceholderProfile | None,
        figure_caption_template: FigureCaptionTemplate | None,
        markdown_table_template: MarkdownTableTemplate | None,
    ) -> None:
        row_count = len(table_block.rows) + (1 if table_block.headers else 0)
        if row_count == 0:
            raise UnsupportedMarkdownError("Tabela vazia nao e suportada na V1.")

        column_count = (
            len(table_block.headers)
            if table_block.headers
            else len(table_block.rows[0])
        )
        table = subdoc.add_table(rows=row_count, cols=column_count)
        self._separate_adjacent_tables(table)

        if markdown_table_template is not None:
            self._apply_markdown_table_template(table, markdown_table_template)
        else:
            try:
                table.style = "Table Grid"
            except (KeyError, ValueError):
                pass

        current_row = 0
        if table_block.headers:
            for column, header_cell in enumerate(table_block.headers):
                self._write_table_cell(
                    table.cell(current_row, column),
                    header_cell,
                    path_imagens=path_imagens,
                    source_dir=source_dir,
                    placeholder_profile=self._resolve_table_cell_profile(
                        markdown_table_template.header_row if markdown_table_template is not None else None,
                        column,
                        fallback_profile=placeholder_profile,
                    ),
                    figure_caption_template=figure_caption_template,
                    cell_template=self._resolve_table_cell_template(
                        markdown_table_template.header_row if markdown_table_template is not None else None,
                        column,
                    ),
                )
            current_row += 1

        for row_index, row in enumerate(table_block.rows):
            row_template = None
            if markdown_table_template is not None:
                row_template = markdown_table_template.odd_row if row_index % 2 == 0 else markdown_table_template.even_row
                self._apply_table_row_template(table.rows[current_row], row_template)
            for column, cell_blocks in enumerate(row):
                self._write_table_cell(
                    table.cell(current_row, column),
                    cell_blocks,
                    path_imagens=path_imagens,
                    source_dir=source_dir,
                    placeholder_profile=self._resolve_table_cell_profile(
                        row_template,
                        column,
                        fallback_profile=placeholder_profile,
                    ),
                    figure_caption_template=figure_caption_template,
                    cell_template=self._resolve_table_cell_template(row_template, column),
                )
            current_row += 1

    @staticmethod
    def _separate_adjacent_tables(table: Any) -> None:
        table_element = table._tbl
        previous_sibling = table_element.getprevious()
        if previous_sibling is None or previous_sibling.tag != qn("w:tbl"):
            return

        separator = OxmlElement("w:p")
        previous_sibling.addnext(separator)

    def _write_table_cell(
        self,
        cell: Any,
        cell_blocks: TableCell,
        *,
        path_imagens: str | None,
        source_dir: str | None,
        placeholder_profile: PlaceholderProfile | None,
        figure_caption_template: FigureCaptionTemplate | None,
        cell_template: TableCellTemplate | None,
    ) -> None:
        paragraph = cell.paragraphs[0]
        self._clear_paragraph_content(paragraph)
        if cell_template is not None:
            self._apply_table_cell_template(cell, cell_template)

        apply_paragraph_profile(
            paragraph,
            placeholder_profile,
            preserve_style=False,
            preserve_indents=False,
        )

        if len(cell_blocks.blocks) == 1 and isinstance(cell_blocks.blocks[0], ParagraphBlock):
            self._add_inline_runs(paragraph, cell_blocks.blocks[0].spans, placeholder_profile=placeholder_profile)
            return

        for block in cell_blocks.blocks:
            self._render_block(
                subdoc=cell,
                block=block,
                path_imagens=path_imagens,
                source_dir=source_dir,
                quote_depth=0,
                placeholder_profile=placeholder_profile,
                figure_caption_template=figure_caption_template,
                quote_block_template=None,
                code_block_template=None,
                markdown_table_template=None,
            )

        if not paragraph.text and len(cell.paragraphs) > 1:
            self._remove_paragraph(paragraph)

    def _render_quote(
        self,
        subdoc: Any,
        quote_block: QuoteBlock,
        *,
        path_imagens: str | None,
        source_dir: str | None,
        quote_depth: int,
        placeholder_profile: PlaceholderProfile | None,
        figure_caption_template: FigureCaptionTemplate | None,
        quote_block_template: BlockContainerTemplate | None,
        code_block_template: BlockContainerTemplate | None,
        markdown_table_template: MarkdownTableTemplate | None,
    ) -> None:
        if quote_block_template is not None:
            self._render_container_blocks(
                parent=subdoc,
                blocks=quote_block.blocks,
                container_template=quote_block_template,
                path_imagens=path_imagens,
                source_dir=source_dir,
                placeholder_profile=quote_block_template.placeholder_profile,
                figure_caption_template=figure_caption_template,
                quote_block_template=quote_block_template,
                code_block_template=code_block_template,
                markdown_table_template=markdown_table_template,
            )
            return

        for block in quote_block.blocks:
            self._render_block(
                subdoc=subdoc,
                block=block,
                path_imagens=path_imagens,
                source_dir=source_dir,
                quote_depth=quote_depth,
                placeholder_profile=placeholder_profile,
                figure_caption_template=figure_caption_template,
                quote_block_template=quote_block_template,
                code_block_template=code_block_template,
                markdown_table_template=markdown_table_template,
            )

    def _render_code_block(
        self,
        subdoc: Any,
        code_block: CodeBlock,
        *,
        quote_depth: int,
        placeholder_profile: PlaceholderProfile | None,
        path_imagens: str | None,
        source_dir: str | None,
        figure_caption_template: FigureCaptionTemplate | None,
        quote_block_template: BlockContainerTemplate | None,
        code_block_template: BlockContainerTemplate | None,
        markdown_table_template: MarkdownTableTemplate | None,
    ) -> None:
        if code_block_template is not None and quote_depth == 0:
            self._render_container_blocks(
                parent=subdoc,
                blocks=[code_block],
                container_template=code_block_template,
                path_imagens=path_imagens,
                source_dir=source_dir,
                placeholder_profile=code_block_template.placeholder_profile,
                figure_caption_template=figure_caption_template,
                quote_block_template=quote_block_template,
                code_block_template=None,
                markdown_table_template=markdown_table_template,
            )
            return

        paragraph = self._add_paragraph(
            subdoc,
            "Quote" if quote_depth > 0 else self._resolve_base_paragraph_style(placeholder_profile),
        )
        apply_paragraph_profile(
            paragraph,
            placeholder_profile,
            preserve_style=quote_depth > 0,
            preserve_indents=quote_depth > 0,
        )
        run = paragraph.add_run(code_block.code)
        apply_run_profile(run, placeholder_profile)
        run.font.name = "Consolas"
        if placeholder_profile is None or placeholder_profile.run.font_size is None:
            run.font.size = Pt(10)

    def _render_container_blocks(
        self,
        *,
        parent: Any,
        blocks: list[BlockLike],
        container_template: BlockContainerTemplate,
        path_imagens: str | None,
        source_dir: str | None,
        placeholder_profile: PlaceholderProfile | None,
        figure_caption_template: FigureCaptionTemplate | None,
        quote_block_template: BlockContainerTemplate | None,
        code_block_template: BlockContainerTemplate | None,
        markdown_table_template: MarkdownTableTemplate | None,
    ) -> None:
        table = self._create_container_table(parent, container_template)
        cell = table.cell(0, 0)
        placeholder_paragraph = cell.paragraphs[0]
        placeholder_paragraph.text = ""

        for block in blocks:
            self._render_block(
                subdoc=cell,
                block=block,
                path_imagens=path_imagens,
                source_dir=source_dir,
                quote_depth=0,
                placeholder_profile=placeholder_profile,
                figure_caption_template=figure_caption_template,
                quote_block_template=quote_block_template,
                code_block_template=code_block_template,
                markdown_table_template=markdown_table_template,
            )

        if not placeholder_paragraph.text and len(cell.paragraphs) > 1:
            self._remove_paragraph(placeholder_paragraph)

    def _apply_markdown_table_template(self, table: Table, template: MarkdownTableTemplate) -> None:
        if template.table_properties_xml is not None:
            self._replace_child(table._tbl, qn("w:tblPr"), deepcopy(template.table_properties_xml))

        if template.grid_column_widths:
            grid_columns = table._tbl.tblGrid.gridCol_lst
            last_width = template.grid_column_widths[-1]
            for index, grid_col in enumerate(grid_columns):
                width = template.grid_column_widths[index] if index < len(template.grid_column_widths) else last_width
                grid_col.w = self._resolve_grid_width_twips(width)

        if table.rows:
            self._apply_table_row_template(table.rows[0], template.header_row)

    def _apply_table_row_template(self, row: Any, row_template: TableRowTemplate) -> None:
        if row_template.row_properties_xml is not None:
            self._replace_child(row._tr, qn("w:trPr"), deepcopy(row_template.row_properties_xml))

        for column_index, cell in enumerate(row.cells):
            cell_template = row_template.resolve_cell(column_index)
            if cell_template is None:
                continue
            self._apply_table_cell_template(cell, cell_template)

    def _apply_table_cell_template(self, cell: Any, cell_template: TableCellTemplate) -> None:
        if cell_template.cell_properties_xml is not None:
            self._replace_child(cell._tc, qn("w:tcPr"), deepcopy(cell_template.cell_properties_xml))

    @staticmethod
    def _resolve_table_cell_template(
        row_template: TableRowTemplate | None,
        column_index: int,
    ) -> TableCellTemplate | None:
        if row_template is None:
            return None
        return row_template.resolve_cell(column_index)

    @staticmethod
    def _resolve_table_cell_profile(
        row_template: TableRowTemplate | None,
        column_index: int,
        *,
        fallback_profile: PlaceholderProfile | None,
    ) -> PlaceholderProfile | None:
        cell_template = row_template.resolve_cell(column_index) if row_template is not None else None
        if cell_template is None or cell_template.paragraph_profile is None:
            return fallback_profile
        return cell_template.paragraph_profile

    def _render_image(
        self,
        subdoc: Any,
        image_block: ImageBlock,
        path_imagens: str | None,
        source_dir: str | None,
        placeholder_profile: PlaceholderProfile | None,
        figure_caption_template: FigureCaptionTemplate | None,
    ) -> None:
        image_path = self._image_resolver.resolve(
            image_block.reference,
            path_imagens=path_imagens,
            source_dir=source_dir,
        )
        has_add_picture = hasattr(subdoc, "add_picture")
        has_add_paragraph = hasattr(subdoc, "add_paragraph")
        has_add_run = hasattr(subdoc, "add_run")

        if has_add_picture:
            inline_shape = subdoc.add_picture(image_path)
            self._resize_image_if_needed(subdoc, inline_shape)
            image_paragraph = subdoc.paragraphs[-1]
            caption_container = subdoc
        elif has_add_paragraph:
            paragraph = subdoc.add_paragraph()
            run = paragraph.add_run()
            inline_shape = run.add_picture(image_path)
            self._resize_image_if_needed(subdoc, inline_shape)
            image_paragraph = paragraph
            caption_container = subdoc
        elif has_add_run:
            run = subdoc.add_run()
            inline_shape = run.add_picture(image_path)
            self._resize_image_if_needed(subdoc, inline_shape)
            image_paragraph = subdoc
            caption_container = getattr(getattr(subdoc, "_element", None), "doc", None)
        else:
            raise TemplateBindingError(
                "O container de renderizacao de imagem nao possui metodo suportado para adicionar imagens.",
                container_type=type(subdoc).__name__,
            )

        self._format_image_paragraph(image_paragraph)

        if image_block.caption and caption_container is not None:
            self._add_figure_caption(
                caption_container,
                image_block.caption,
                placeholder_profile=placeholder_profile,
                figure_caption_template=figure_caption_template,
                anchor_paragraph=image_paragraph,
            )

    def _render_horizontal_rule(
        self,
        subdoc: Any,
        *,
        quote_depth: int,
        placeholder_profile: PlaceholderProfile | None,
    ) -> None:
        paragraph = self._add_paragraph(
            subdoc,
            "Quote" if quote_depth > 0 else self._resolve_base_paragraph_style(placeholder_profile),
        )
        apply_paragraph_profile(
            paragraph,
            placeholder_profile,
            preserve_style=quote_depth > 0,
            preserve_indents=quote_depth > 0,
        )
        paragraph_properties = paragraph._p.get_or_add_pPr()

        existing_borders = paragraph_properties.find(qn("w:pBdr"))
        if existing_borders is not None:
            paragraph_properties.remove(existing_borders)

        paragraph_borders = OxmlElement("w:pBdr")
        bottom_border = OxmlElement("w:bottom")
        bottom_border.set(qn("w:val"), "single")
        bottom_border.set(qn("w:sz"), "6")
        bottom_border.set(qn("w:space"), "1")
        bottom_border.set(qn("w:color"), "auto")
        paragraph_borders.append(bottom_border)
        paragraph_properties.append(paragraph_borders)

    def _render_spacer(
        self,
        subdoc: Any,
        *,
        line_breaks: int,
        quote_depth: int,
        placeholder_profile: PlaceholderProfile | None,
    ) -> None:
        for _ in range(max(0, line_breaks)):
            paragraph = self._add_paragraph(
                subdoc,
                "Quote" if quote_depth > 0 else self._resolve_base_paragraph_style(placeholder_profile),
            )
            apply_paragraph_profile(
                paragraph,
                placeholder_profile,
                preserve_style=quote_depth > 0,
                preserve_indents=quote_depth > 0,
            )

    def _add_figure_caption(
        self,
        subdoc: Any,
        caption: str,
        *,
        placeholder_profile: PlaceholderProfile | None,
        figure_caption_template: FigureCaptionTemplate | None,
        anchor_paragraph: Any | None = None,
    ) -> None:
        self._figure_sequence += 1
        if figure_caption_template is not None:
            self._render_figure_caption_from_template(
                figure_caption_template,
                caption,
                current_value=self._figure_sequence,
                anchor_paragraph=anchor_paragraph,
            )
            return

        caption_paragraph = self._add_paragraph(subdoc, "Caption")
        caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        figure_run = caption_paragraph.add_run("Figura ")
        self._append_sequence_field(
            caption_paragraph,
            sequence_name="Figura",
            current_value=self._figure_sequence,
            placeholder_profile=None,
        )
        caption_run = caption_paragraph.add_run(f" - {caption}")
        apply_run_profile(caption_run, None)

    def _resize_image_if_needed(self, subdoc: Any, inline_shape: Any) -> None:
        max_width = self._get_available_page_width(subdoc)
        max_height = self._get_available_page_height(subdoc)
        current_width = int(inline_shape.width)
        current_height = int(inline_shape.height)

        if current_width <= int(max_width) and current_height <= int(max_height):
            return

        width_ratio = current_width / int(max_width)
        height_ratio = current_height / int(max_height)
        size_ratio = max(width_ratio, height_ratio)
        reduction_scale = self._resolve_reduction_scale(size_ratio)
        fit_scale = min(
            int(max_width) / current_width,
            int(max_height) / current_height,
        )
        final_scale = min(reduction_scale, fit_scale)

        inline_shape.width = Emu(max(1, int(current_width * final_scale)))
        inline_shape.height = Emu(max(1, int(current_height * final_scale)))

    @staticmethod
    def _resolve_reduction_scale(size_ratio: float) -> float:
        if size_ratio <= 1.5:
            return 0.6
        if size_ratio <= 2.5:
            return 0.5
        if size_ratio <= 3.5:
            return 0.4
        return 0.3

    @staticmethod
    def _format_image_paragraph(paragraph: Any) -> None:
        paragraph_element = getattr(paragraph, "_p", None)
        if paragraph_element is None:
            return

        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

        paragraph_format = paragraph.paragraph_format
        paragraph_format.left_indent = 0
        paragraph_format.right_indent = 0
        paragraph_format.first_line_indent = 0

    @staticmethod
    def _get_available_page_width(subdoc: Any) -> Emu:
        if not hasattr(subdoc, "sections") or not subdoc.sections:
            return Emu(int(Inches(6.5)))
        section = subdoc.sections[-1]
        page_width = int(section.page_width or Inches(8.5))
        left_margin = int(section.left_margin or Inches(1))
        right_margin = int(section.right_margin or Inches(1))
        return Emu(page_width - left_margin - right_margin)

    @staticmethod
    def _get_available_page_height(subdoc: Any) -> Emu:
        if not hasattr(subdoc, "sections") or not subdoc.sections:
            return Emu(int(Inches(9)))
        section = subdoc.sections[-1]
        page_height = int(section.page_height or Inches(11))
        top_margin = int(section.top_margin or Inches(1))
        bottom_margin = int(section.bottom_margin or Inches(1))
        return Emu(page_height - top_margin - bottom_margin)

    def _append_sequence_field(
        self,
        paragraph: Any,
        *,
        sequence_name: str,
        current_value: int,
        placeholder_profile: PlaceholderProfile | None,
    ) -> None:
        begin_run = paragraph.add_run()
        apply_run_profile(begin_run, placeholder_profile)
        begin = OxmlElement("w:fldChar")
        begin.set(qn("w:fldCharType"), "begin")
        begin_run._r.append(begin)

        instruction_run = paragraph.add_run()
        apply_run_profile(instruction_run, placeholder_profile)
        instruction = OxmlElement("w:instrText")
        instruction.set(qn("xml:space"), "preserve")
        instruction.text = f" SEQ {sequence_name} \\* ARABIC "
        instruction_run._r.append(instruction)

        separate_run = paragraph.add_run()
        apply_run_profile(separate_run, placeholder_profile)
        separate = OxmlElement("w:fldChar")
        separate.set(qn("w:fldCharType"), "separate")
        separate_run._r.append(separate)

        value_run = paragraph.add_run(str(current_value))
        apply_run_profile(value_run, placeholder_profile)

        end_run = paragraph.add_run()
        apply_run_profile(end_run, placeholder_profile)
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        end_run._r.append(end)

    def _add_inline_runs(
        self,
        paragraph: Any,
        spans: list[InlineSpan],
        *,
        placeholder_profile: PlaceholderProfile | None,
    ) -> None:
        for span in spans:
            if isinstance(span, LinkSpan):
                self._add_hyperlink(paragraph, span, placeholder_profile=placeholder_profile)
                continue

            run = paragraph.add_run()
            self._populate_run(run, span, placeholder_profile=placeholder_profile)

    def _add_hyperlink(
        self,
        paragraph: Any,
        link_span: LinkSpan,
        *,
        placeholder_profile: PlaceholderProfile | None,
    ) -> None:
        if not link_span.target or link_span.target.startswith("#"):
            self._add_inline_runs(paragraph, link_span.spans, placeholder_profile=placeholder_profile)
            return

        relationship_id = paragraph.part.relate_to(link_span.target, RT.HYPERLINK, is_external=True)
        hyperlink = OxmlElement("w:hyperlink")
        hyperlink.set(qn("r:id"), relationship_id)
        paragraph._p.append(hyperlink)

        for span in link_span.spans:
            if isinstance(span, LinkSpan):
                raise UnsupportedMarkdownError(
                    "Links aninhados nao fazem parte da V1.",
                    span_type=type(span).__name__,
                )

            run_element = OxmlElement("w:r")
            hyperlink.append(run_element)
            run = Run(run_element, paragraph)
            self._populate_run(
                run,
                span,
                within_hyperlink=True,
                placeholder_profile=placeholder_profile,
            )

    def _populate_run(
        self,
        run: Run,
        span: InlineSpan,
        *,
        within_hyperlink: bool = False,
        placeholder_profile: PlaceholderProfile | None,
    ) -> None:
        if isinstance(span, LinkSpan):
            raise UnsupportedMarkdownError(
                "Span de link inesperado ao popular run.",
                span_type=type(span).__name__,
            )

        run.text = span.text
        apply_run_profile(run, placeholder_profile)

        if within_hyperlink:
            self._apply_hyperlink_style(run)

        if isinstance(span, BoldItalicSpan):
            run.bold = True
            run.italic = True
            return

        if isinstance(span, BoldSpan):
            run.bold = True
            return

        if isinstance(span, ItalicSpan):
            run.italic = True
            return

        if isinstance(span, CodeSpan):
            run.font.name = "Consolas"
            if placeholder_profile is None or placeholder_profile.run.font_size is None:
                run.font.size = Pt(10)
            return

        if isinstance(span, TextSpan):
            return

        raise UnsupportedMarkdownError(
            "Span inline fora da cobertura da V1.",
            span_type=type(span).__name__,
        )

    @staticmethod
    def _apply_hyperlink_style(run: Run) -> None:
        try:
            run.style = "Hyperlink"
        except (KeyError, ValueError):
            run.font.color.rgb = RGBColor(0x05, 0x63, 0xC1)
            run.underline = True

    @staticmethod
    def _resolve_base_paragraph_style(profile: PlaceholderProfile | None) -> str | None:
        if profile is None:
            return None
        return profile.paragraph.style_name

    def _render_figure_caption_from_template(
        self,
        figure_caption_template: FigureCaptionTemplate,
        caption: str,
        *,
        current_value: int,
        anchor_paragraph: Any | None,
    ) -> None:
        paragraph_xml = deepcopy(figure_caption_template.paragraph_xml)
        self._replace_figure_caption_marker_in_xml(paragraph_xml, caption)
        self._update_first_sequence_display_value_in_xml(paragraph_xml, current_value)

        if anchor_paragraph is None:
            raise UnsupportedMarkdownError(
                "Nao foi possivel determinar a ancora da legenda de figura.",
                placeholder="{{ figure_caption }}",
            )

        anchor_paragraph._p.addnext(paragraph_xml)

    def _replace_figure_caption_marker_in_xml(self, paragraph_xml: Any, caption: str) -> None:
        text_elements = list(paragraph_xml.iter(qn("w:t")))
        if not text_elements:
            raise UnsupportedMarkdownError(
                "Paragrafo de legenda de figura no template nao possui texto.",
                placeholder="{{ figure_caption }}",
            )

        combined_text = "".join(element.text or "" for element in text_elements)
        match = FIGURE_CAPTION_PATTERN.search(combined_text)
        if match is None:
            raise UnsupportedMarkdownError(
                "Paragrafo de legenda de figura no template nao contem o placeholder esperado.",
                placeholder="{{ figure_caption }}",
            )

        start, end = match.span()
        position = 0
        caption_written = False
        for element in text_elements:
            element_text = element.text or ""
            element_start = position
            element_end = position + len(element_text)
            position = element_end

            if element_end <= start or element_start >= end:
                continue

            prefix = element_text[: max(0, start - element_start)]
            suffix = element_text[min(len(element_text), end - element_start) :]

            if caption_written:
                element.text = suffix
                continue

            element.text = f"{prefix}{caption}{suffix}"
            caption_written = True

    def _update_first_sequence_display_value_in_xml(self, paragraph_xml: Any, current_value: int) -> None:
        inside_sequence_field = False
        awaiting_result = False

        for run_element in paragraph_xml.iter(qn("w:r")):
            for instruction in run_element.findall(qn("w:instrText")):
                if "SEQ" in (instruction.text or ""):
                    inside_sequence_field = True

            for field_char in run_element.findall(qn("w:fldChar")):
                field_type = field_char.get(qn("w:fldCharType"))
                if field_type == "separate" and inside_sequence_field:
                    awaiting_result = True
                    continue
                if field_type == "end" and inside_sequence_field:
                    return

            if awaiting_result:
                text_elements = run_element.findall(qn("w:t"))
                if text_elements:
                    text_elements[0].text = str(current_value)
                    for extra_text in text_elements[1:]:
                        extra_text.text = ""
                    return

    def _add_paragraph(
        self,
        subdoc: Any,
        style: str | None = None,
        *,
        fallback_style: str | None = None,
    ) -> Any:
        if style is None:
            return subdoc.add_paragraph()

        try:
            return subdoc.add_paragraph(style=style)
        except (KeyError, ValueError):
            if fallback_style is not None and fallback_style != style:
                try:
                    return subdoc.add_paragraph(style=fallback_style)
                except (KeyError, ValueError):
                    pass
            return subdoc.add_paragraph()

    def _create_container_table(self, parent: Any, container_template: BlockContainerTemplate) -> Table:
        try:
            table = parent.add_table(rows=1, cols=1)
        except Exception as exc:
            raise TemplateBindingError(
                "Nao foi possivel criar a tabela-prototipo do container de renderizacao.",
                container_type=type(parent).__name__,
            ) from exc

        if container_template.table_properties_xml is not None:
            self._replace_child(table._tbl, qn("w:tblPr"), deepcopy(container_template.table_properties_xml))

        if container_template.grid_column_widths:
            grid_columns = table._tbl.tblGrid.gridCol_lst
            last_width = container_template.grid_column_widths[-1]
            for index, grid_col in enumerate(grid_columns):
                width = (
                    container_template.grid_column_widths[index]
                    if index < len(container_template.grid_column_widths)
                    else last_width
                )
                grid_col.w = self._resolve_grid_width_twips(width)

        if container_template.row_properties_xml is not None:
            self._replace_child(table.rows[0]._tr, qn("w:trPr"), deepcopy(container_template.row_properties_xml))

        if container_template.cell_properties_xml is not None:
            self._replace_child(table.cell(0, 0)._tc, qn("w:tcPr"), deepcopy(container_template.cell_properties_xml))

        return table

    @staticmethod
    def _remove_paragraph(paragraph: Any) -> None:
        element = paragraph._element
        parent = element.getparent()
        if parent is not None:
            parent.remove(element)

    @staticmethod
    def _clear_paragraph_content(paragraph: Any) -> None:
        for child in list(paragraph._p):
            if child.tag != qn("w:pPr"):
                paragraph._p.remove(child)

    @staticmethod
    def _replace_child(parent: Any, tag: str, new_child: Any) -> None:
        existing_child = parent.find(tag)
        if existing_child is not None:
            parent.remove(existing_child)
        parent.insert(0, new_child)

    @staticmethod
    def _resolve_grid_width_twips(width: Any) -> Any:
        return width.twips if hasattr(width, "twips") else deepcopy(width)
